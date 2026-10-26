#!/usr/bin/env python3
# -*- coding: utf-8 -*-
# /// script
# requires-python = ">=3.9"
# dependencies = ["openpyxl>=3.1"]
# ///
"""
VCB Batch Transfer Payroll Converter (Converter Bảng Lương sang File Lô VCB)
Converts raw payroll data (XLSX, CSV, JSON) to 100% exact Vietcombank Batch Transfer Format (File Lô Mẫu VCB).
"""

import sys
import os
import csv
import json
import re
import argparse
import unicodedata
from pathlib import Path

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:  # dry-run with JSON/CSV input works without openpyxl
    openpyxl = None

OPENPYXL_HINT = "Thiếu openpyxl. Cài: pip install openpyxl  (hoặc chạy: uv run scripts/convert_payroll.py ...)"

# Load Bank Mapping Database
SCRIPT_DIR = Path(__file__).resolve().parent
RESOURCES_DIR = SCRIPT_DIR.parent / "resources"
MAPPING_FILE = RESOURCES_DIR / "bank_mapping.json"

def remove_vietnamese_accent(text: str) -> str:
    if not text:
        return ""
    text = unicodedata.normalize('NFD', text)
    text = ''.join(c for c in text if unicodedata.category(c) != 'Mn')
    text = text.replace('đ', 'd').replace('Đ', 'D')
    return text

def remove_vietnamese_accent_upper(text: str) -> str:
    return remove_vietnamese_accent(text).upper().strip()

def load_bank_mappings():
    if MAPPING_FILE.exists():
        with open(MAPPING_FILE, 'r', encoding='utf-8') as f:
            return json.load(f)
    return []

BANK_DATABASE = load_bank_mappings()

def resolve_bank_code_and_type(bank_raw: str) -> tuple[str, str]:
    """
    Returns (transfer_type_str, citad_code_or_vcb_string)
    Type 1: VCB Internal -> ("1", "Vietcombank")
    Type 2: External CITAD -> ("2", "79303001" / "01311001" / "01358001"...)
    VCB Digibank batch parser REQUIRES the 8-digit CITAD Bank Code for type 2 transfers!
    """
    clean_bank = remove_vietnamese_accent_upper(bank_raw)
    
    # Check if Vietcombank
    if any(k in clean_bank for k in ["VIETCOMBANK", "VCB", "NGOAI THUONG"]):
        return "1", "Vietcombank"
    
    # Check Bank Database Aliases (Sorted by alias length descending with word boundaries)
    for bank in BANK_DATABASE:
        for alias in sorted(bank.get("aliases", []), key=len, reverse=True):
            clean_alias = remove_vietnamese_accent_upper(alias)
            if clean_alias:
                pattern = r'(?:\b|_)' + re.escape(clean_alias) + r'(?:\b|_)'
                if re.search(pattern, clean_bank):
                    return "2", str(bank["code"])
                    
    # Check Bank Database Full Names
    for bank in BANK_DATABASE:
        clean_name = remove_vietnamese_accent_upper(bank["name"])
        if clean_name and clean_name in clean_bank:
            return "2", str(bank["code"])
            
    # Fallback to cleaned raw text if not in database
    return "2", bank_raw.strip()

def convert_payroll_to_vcb(
    input_records: list[dict],
    output_xlsx_path: str,
    debit_account: str = "1052233438",
    default_fee: str = "B "
):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    
    # Header Columns Definition (Exact VCB File Lo Mau Spec)
    headers = [
        "STT",
        "CUS REF\n( KH tùy chọn)",
        "Loại hình chuyển tiền \n(Select)",
        "Tài khoản trích nợ\n(Text)",
        "Tài khoản người hưởng\n(Text)",
        "Tên người hưởng\n(Text)",
        "Địa chỉ người hưởng\n(Text)",
        "Số CMT\n(Text)",
        "Ngày cấp\n(Text)",
        "Nơi cấp\n(Text)",
        "Tên Ngân hàng hưởng\n(Text)",
        "Số tiền\n(số)",
        "Loại tiền chuyển\n(Text)",
        "Phí\n(Select)",
        "Nội dung lệnh\n(Text)"
    ]
    
    # Header Styling
    header_fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
    header_font = Font(name="Calibri", size=10, bold=True)
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style='thin', color='BFBFBF'),
        right=Side(style='thin', color='BFBFBF'),
        top=Side(style='thin', color='BFBFBF'),
        bottom=Side(style='thin', color='BFBFBF')
    )
    
    data_font = Font(name="Calibri", size=10)
    
    # Write Header Row
    ws.row_dimensions[1].height = 40
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_align
        cell.border = thin_border
        
    start_cus_ref = 2000000014
    
    # Write Data Rows
    for idx, rec in enumerate(input_records, 1):
        row_num = idx + 1
        ws.row_dimensions[row_num].height = 20
        
        stt = idx
        cus_ref = rec.get("cus_ref") or str(start_cus_ref + idx - 1)
        acc_no = str(rec.get("account_no", "")).strip()
        beneficiary_name_raw = str(rec.get("name", "")).strip()
        beneficiary_name = remove_vietnamese_accent_upper(beneficiary_name_raw)
        bank_raw = str(rec.get("bank", "")).strip()
        
        # Resolve CITAD Bank Code or Vietcombank
        transfer_type, bank_target_code = resolve_bank_code_and_type(bank_raw)
        
        amount_val = parse_amount(rec.get("amount", 0))

        base_desc_raw = str(rec.get("description", "Thanh toan luong")).strip()
        base_desc_no_accent = remove_vietnamese_accent(base_desc_raw)
        
        # Rule: Append ' - [Tên nhân viên KHÔNG DẤU]' if not already present
        if beneficiary_name and not base_desc_no_accent.endswith(f"- {beneficiary_name}"):
            description_str = f"{base_desc_no_accent} - {beneficiary_name}"
        else:
            description_str = base_desc_no_accent
            
        # Ensure 100% NO ACCENTS in description
        description = remove_vietnamese_accent(description_str)
        
        row_values = [
            stt,
            str(cus_ref),
            int(transfer_type),
            str(debit_account),
            str(acc_no),
            beneficiary_name,
            "", # Địa chỉ người hưởng
            "", # Số CMT
            "", # Ngày cấp
            "", # Nơi cấp
            bank_target_code, # Mã CITAD 8 số (79303001, 01311001...) hoặc Vietcombank
            amount_val,
            "VND",
            default_fee,
            description
        ]
        
        for col_idx, val in enumerate(row_values, 1):
            cell = ws.cell(row=row_num, column=col_idx, value=val)
            cell.font = data_font
            cell.border = thin_border
            
            # Formats & Alignments
            if col_idx in [1, 3]: # STT, Loại hình
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx in [2, 4, 5, 11]: # Text IDs, Account numbers & CITAD Bank Code
                cell.number_format = '@'
                cell.alignment = Alignment(horizontal="left", vertical="center")
            elif col_idx == 12: # Amount
                cell.number_format = '#,##0'
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif col_idx in [13, 14]: # VND, Phí
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
                
    # Auto-adjust column widths
    col_widths = {
        1: 6,   # STT
        2: 16,  # CUS REF
        3: 20,  # Loại hình
        4: 20,  # TK trích nợ
        5: 22,  # TK người hưởng
        6: 25,  # Tên người hưởng
        7: 20,  # Địa chỉ
        8: 15,  # CMT
        9: 12,  # Ngày cấp
        10: 15, # Nơi cấp
        11: 22, # Tên Ngân hàng hưởng (Mã CITAD)
        12: 18, # Số tiền
        13: 15, # Loại tiền
        14: 10, # Phí
        15: 45  # Nội dung
    }
    
    for col_idx, width in col_widths.items():
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = width
        
    wb.save(output_xlsx_path)
    print(f"[SUCCESS] Exported VCB Batch File to: {output_xlsx_path}")

def load_records_from_file(file_path: str) -> list[dict]:
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"Input file not found: {file_path}")
        
    if path.suffix.lower() == '.json':
        with open(path, 'r', encoding='utf-8') as f:
            return json.load(f)
    elif path.suffix.lower() == '.csv':
        # Read every cell as text so leading zeros in account numbers survive
        with open(path, 'r', encoding='utf-8-sig', newline='') as f:
            return [
                {str(k).strip().lower(): v for k, v in row.items() if k}
                for row in csv.DictReader(f)
                if any((v or '').strip() for v in row.values())
            ]
    elif path.suffix.lower() == '.xlsx':
        if openpyxl is None:
            raise RuntimeError(OPENPYXL_HINT)
        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(h).strip().lower() if h else '' for h in rows[0]]
        records = []
        for r in rows[1:]:
            if not any(r):
                continue
            rec = {}
            for idx, val in enumerate(r):
                if idx < len(headers):
                    rec[headers[idx]] = val
            records.append(rec)
        return records
    else:
        raise ValueError(f"Unsupported file format: {path.suffix}")

def parse_amount(raw) -> float:
    """Parse a VND amount; shared by convert_payroll_to_vcb and summarize so preview == file.

    Accepts 7500000, "7.500.000", "7,500,000", "7 500 000 đ", "7.500.000,50", "7500000.0".
    Anything ambiguous returns 0.0, which summarize() reports as a warning.
    """
    if isinstance(raw, (int, float)):
        return float(raw)
    s = re.sub(r'[^0-9.,]', '', str(raw))
    if re.fullmatch(r'\d+', s):
        return float(s)
    for sep, dec in (('.', ','), (',', '.')):
        m = re.fullmatch(rf'(\d{{1,3}}(?:\{sep}\d{{3}})+)(?:\{dec}(\d+))?', s)
        if m:
            return float(m.group(1).replace(sep, '') + ('.' + m.group(2) if m.group(2) else ''))
    m = re.fullmatch(r'(\d+)[.,](\d{1,2})', s)
    if m:
        return float(f"{m.group(1)}.{m.group(2)}")
    return 0.0


def summarize(records: list[dict]) -> dict:
    """Preview what convert_payroll_to_vcb will write, plus warnings a human must clear."""
    rows, warnings = [], []
    for idx, rec in enumerate(records, 1):
        raw_amount = rec.get("amount", 0)
        amount = parse_amount(raw_amount)
        acc = str(rec.get("account_no", "")).strip()
        bank_raw = str(rec.get("bank", "")).strip()
        transfer_type, bank_code = resolve_bank_code_and_type(bank_raw)
        name = remove_vietnamese_accent_upper(str(rec.get("name", "")))
        rows.append({"stt": idx, "name": name, "account_no": acc, "transfer_type": int(transfer_type),
                     "bank": bank_code, "amount": amount})
        if amount <= 0:
            warnings.append(f"Dòng {idx}: số tiền '{raw_amount}' đọc ra {amount:g}")
        elif amount != int(amount):
            warnings.append(f"Dòng {idx}: số tiền '{raw_amount}' có phần lẻ ({amount:g}) — VND không có xu")
        if not acc.isdigit():
            warnings.append(f"Dòng {idx}: số tài khoản '{acc}' không chỉ gồm chữ số")
        if not name:
            warnings.append(f"Dòng {idx}: thiếu tên người hưởng")
        if transfer_type == "2" and not re.fullmatch(r'\d{8}', bank_code):
            warnings.append(f"Dòng {idx}: không tìm được mã Citad cho ngân hàng '{bank_raw}'")
    return {"rows": len(rows), "total_amount": sum(r["amount"] for r in rows),
            "warnings": warnings, "records": rows}


def main(argv=None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(
        description="Chuyển bảng lương (JSON/CSV/XLSX) sang file lô chuyển tiền Vietcombank (.xlsx).",
        epilog="Cột đầu vào: name, account_no, bank, amount, description (tùy chọn), cus_ref (tùy chọn).\n"
               "Ví dụ: python scripts/convert_payroll.py --input luong.csv --debit-account 0123456789 --dry-run\n"
               "Exit code: 0 OK, 1 lỗi đọc/ghi, 2 có cảnh báo (dùng --allow-warnings để vẫn xuất), 3 thiếu openpyxl.",
        formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("legacy", nargs="*", help=argparse.SUPPRESS)  # old form: <input> [output] [debit_account]
    parser.add_argument("--input", help="File bảng lương: .json, .csv hoặc .xlsx")
    parser.add_argument("--output", default=None, help="File .xlsx xuất ra (mặc định: File_Lo_VCB_Out.xlsx)")
    parser.add_argument("--debit-account", default=None, help="Số tài khoản công ty trích nợ (bắt buộc khi xuất file)")
    parser.add_argument("--dry-run", action="store_true", help="Chỉ in tóm tắt JSON (số dòng, tổng tiền, cảnh báo), không ghi file")
    parser.add_argument("--allow-warnings", action="store_true", help="Vẫn xuất file khi còn cảnh báo (chỉ dùng sau khi người dùng đã xem)")
    args = parser.parse_args(argv)

    legacy = args.legacy
    input_file = args.input or (legacy[0] if legacy else None)
    output_file = args.output or (legacy[1] if len(legacy) > 1 else "File_Lo_VCB_Out.xlsx")
    debit_acc = args.debit_account or (legacy[2] if len(legacy) > 2 else None)
    if not input_file:
        parser.error("thiếu --input")

    try:
        records = load_records_from_file(input_file)
    except RuntimeError as e:
        print(f"[ERROR] {e}", file=sys.stderr)
        return 3
    except Exception as e:
        print(f"[ERROR] Không đọc được file bảng lương: {e}", file=sys.stderr)
        return 1

    summary = summarize(records)
    if args.dry_run:
        print(json.dumps(summary, ensure_ascii=False, indent=2))
        return 2 if summary["warnings"] else 0

    if not debit_acc:
        parser.error("thiếu --debit-account (tài khoản công ty trích nợ); không dùng giá trị mặc định cho lệnh chi tiền")
    if summary["warnings"] and not args.allow_warnings:
        print(json.dumps({"warnings": summary["warnings"]}, ensure_ascii=False, indent=2), file=sys.stderr)
        print("[STOP] Còn cảnh báo — sửa dữ liệu nguồn hoặc chạy lại với --allow-warnings sau khi người dùng xác nhận.", file=sys.stderr)
        return 2
    if openpyxl is None:
        print(f"[ERROR] {OPENPYXL_HINT}", file=sys.stderr)
        return 3
    try:
        convert_payroll_to_vcb(records, output_file, debit_acc)
    except Exception as e:
        print(f"[ERROR] Failed to convert payroll file: {e}", file=sys.stderr)
        return 1
    print(json.dumps({"output": output_file, "rows": summary["rows"], "total_amount": summary["total_amount"]},
                     ensure_ascii=False))
    return 0


if __name__ == "__main__":
    sys.exit(main())
